"""Reads documents off disk into sources.

Mixed formats are handled by extracting text and nothing else. Layout, styling
and images are discarded deliberately: every claim is later located by character
offset in this text, so the text a citation points into has to be stable and has
to be the same text the model was shown.

A format that is not supported is refused by name rather than read as bytes and
half-understood.
"""

from __future__ import annotations

from pathlib import Path

SUPPORTED_SUFFIXES = frozenset({".md", ".txt", ".docx", ".pdf"})


class UnsupportedFormat(Exception):
    """Raised for a file this system does not claim to read."""


def read_text(path: Path) -> str:
    """Extract the text of one document."""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise UnsupportedFormat(
            f"{path.name} has extension {suffix or '(none)'}, which this system does not read. "
            f"Supported formats: {supported}."
        )

    if suffix in (".md", ".txt"):
        return path.read_text(encoding="utf-8")
    if suffix == ".docx":
        return _read_docx(path)
    return _read_pdf(path)


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(blocks)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def discover(folder: Path) -> tuple[Path, ...]:
    """Every readable document in a folder, in a stable order.

    Sorted so that two runs over the same folder process the same documents in
    the same order, which is what makes a rerun comparable to the run before it.
    """
    if not folder.exists():
        return ()
    return tuple(
        sorted(
            p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
        )
    )
